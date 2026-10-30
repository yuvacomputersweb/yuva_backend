"""
Import past sales from the shop's sales register (Excel).

Reads the 'Sales Register' sheet (or the first sheet) by column name, so
columns can be in any order. Recognised headers (case-insensitive):
  Sale Date / Date, Customer Name / Name, Mobile / Phone, Store,
  Laptop / Model / Model, Amount, Notes, S.No
Re-importing the same file updates rows instead of duplicating them.
"""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl

from .models import CustomerPurchase

HEADERS = {
    'date': ('sale date', 'date', 'purchase date'),
    'name': ('customer name', 'name', 'customer'),
    'phone': ('mobile', 'phone', 'mobile number', 'phone number', 'whatsapp number'),
    'store': ('store', 'branch'),
    'model': ('laptop / model', 'model', 'laptop', 'product'),
    'amount': ('amount (₹)', 'amount', 'price', 'sale amount'),
    'notes': ('notes', 'remarks'),
    'ref': ('s.no', 's no', 'sno', 'sl no', 'id'),
}


def _phone(value) -> str | None:
    digits = re.sub(r'\D', '', str(value or ''))
    if digits.endswith('.0'):
        digits = digits[:-2]
    digits = digits[-10:]
    return digits if len(digits) == 10 and digits[0] in '6789' else None


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d', '%d.%m.%Y'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except (TypeError, ValueError):
            continue
    return None


def _amount(value):
    try:
        return Decimal(str(value).replace(',', '').replace('₹', '').strip()) if value not in (None, '') else None
    except InvalidOperation:
        return None


def import_sales_register(file, default_store: str = 'Dilsukhnagar HQ', batch_size: int = 2000) -> dict:
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb['Sales Register'] if 'Sales Register' in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or '').strip().lower() for h in next(rows, [])]
    col = {key: next((header.index(n) for n in names if n in header), None) for key, names in HEADERS.items()}
    if col['date'] is None or col['phone'] is None:
        raise ValueError("The sheet needs at least a 'Sale Date' and a 'Mobile' column.")

    get = lambda row, key: row[col[key]] if col[key] is not None and col[key] < len(row) else None
    stats = {'imported': 0, 'skipped_no_phone': 0, 'skipped_no_date': 0, 'rows': 0}
    batch = []

    def flush():
        CustomerPurchase.objects.bulk_create(
            batch, update_conflicts=True, unique_fields=['source_ref'],
            update_fields=['name', 'phone', 'sale_date', 'store', 'model_name', 'amount', 'notes'],
        )
        stats['imported'] += len(batch)
        batch.clear()

    for i, row in enumerate(rows, start=2):
        if not any(v not in (None, '') for v in row):
            continue
        stats['rows'] += 1
        phone = _phone(get(row, 'phone'))
        if not phone:
            stats['skipped_no_phone'] += 1
            continue
        sale_date = _date(get(row, 'date'))
        if not sale_date:
            stats['skipped_no_date'] += 1
            continue
        ref = get(row, 'ref')
        batch.append(CustomerPurchase(
            name=str(get(row, 'name') or '').strip()[:150],
            phone=phone,
            sale_date=sale_date,
            store=str(get(row, 'store') or default_store).strip()[:80],
            model_name=str(get(row, 'model') or '').strip()[:200],
            amount=_amount(get(row, 'amount')),
            notes=str(get(row, 'notes') or '').strip(),
            source_ref=f"{ws.title[:20]}:{int(ref) if isinstance(ref, (int, float)) else f'r{i}'}",
        ))
        if len(batch) >= batch_size:
            flush()
    if batch:
        flush()
    return stats
